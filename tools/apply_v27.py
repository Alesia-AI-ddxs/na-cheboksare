# -*- coding: utf-8 -*-
"""v2.6 -> v2.7: языковой модуль, этап 1 «Каркас» (ТЗ мультиязычности v1.3).

Что делает (старую v2.6 НЕ трогает — работает с копией):
 1) листы ЯЗЫКИ / ИМЕНА / ВЫЧИТКА переносятся из таблицы-дополнения (со стилями);
    в ЯЗЫКИ добавляется колонка «Тест-строка (тофу)»; старый лист «Языки» (требования к дикторам)
    переименован в «Языки — дикторы» — Excel не допускает два листа «Языки» и «ЯЗЫКИ» (регистр не различает);
 2) код татарского приводится к одному: tt-RU -> tt (в шапках листов); слово «мордовский» из шапок убрано;
 3) в листах «Интерфейс», «Тексты», «Легенды» добавлены колонки всех письменных вариантов реестра,
    которых ещё не было; ar (арабский, MSA) заполнен ЧЕРНОВИКОМ перевода ИИ, остальные — пустые (пусто = фолбэк);
 4) новые ключи интерфейса языкового модуля (8 готовых языков + ar);
 5) журнал v2.6 -> v2.7 на листе «Как устроено».
CSV Unity не меняются (их структура закреплена контрактом; новые ключи — только для приложения).

Запуск (из папки app_build):  python3 -I apply_v27.py
"""
import shutil
import sys
from copy import copy
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment
from openpyxl.utils import get_column_letter

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))
from texts_v27 import AR_LEGENDS, AR_TEXTS, AR_UI, L9, NEW_UI, TOFU_TEST  # noqa: E402

SRC = HERE.parent / 'xlsx_work' / 'КОНТЕНТ_На_Чебоксаре_v2.6_2026-10-22.xlsx'
DST = HERE.parent / 'xlsx_work' / 'КОНТЕНТ_На_Чебоксаре_v2.7_2026-10-22.xlsx'
ADD = HERE.parent / 'lang_module' / 'module' / 'МАСТЕР_ТАБЛИЦА_ДОПОЛНЕНИЕ_ЯЗЫКИ_v1.0_2026-10-22.xlsx'

shutil.copyfile(SRC, DST)
wb = openpyxl.load_workbook(DST, data_only=False)
add = openpyxl.load_workbook(ADD, data_only=False)
log = []


def copy_style(src, dst):
    if src.has_style:
        dst.font = copy(src.font)
        dst.fill = copy(src.fill)
        dst.border = copy(src.border)
        dst.alignment = copy(src.alignment)
        dst.number_format = src.number_format
        dst.protection = copy(src.protection)


# ------------------------------------------------------------------ 1. старый лист «Языки» -> «Языки — дикторы»
old = wb['Языки']
old.title = 'Языки — дикторы'
for r in range(2, old.max_row + 1):
    if old.cell(r, 1).value == 'tt-RU':
        old.cell(r, 1).value = 'tt'
    v = old.cell(r, 2).value
    if isinstance(v, str) and '(мордовский)' in v:
        old.cell(r, 2).value = v.replace(' (мордовский)', '')
log.append('Лист «Языки» (требования к дикторам) переименован в «Языки — дикторы»: Excel не различает регистр, '
           'а новый реестр называется «ЯЗЫКИ». Код tt-RU -> tt; «(мордовский)» в названиях убрано.')

# ------------------------------------------------------------------ 2. листы из дополнения
for name in ['ЯЗЫКИ', 'ИМЕНА', 'ВЫЧИТКА']:
    s = add[name]
    d = wb.create_sheet(name)
    for row in s.iter_rows():
        for c in row:
            n = d.cell(c.row, c.column, c.value)
            copy_style(c, n)
    for k, dim in s.column_dimensions.items():
        d.column_dimensions[k].width = dim.width
    for k, dim in s.row_dimensions.items():
        if dim.height:
            d.row_dimensions[k].height = dim.height
    d.freeze_panes = s.freeze_panes
    for m in s.merged_cells.ranges:
        d.merge_cells(str(m))
    if s.sheet_properties.tabColor is not None:
        d.sheet_properties.tabColor = copy(s.sheet_properties.tabColor)

ws_lang = wb['ЯЗЫКИ']
# колонка «Тест-строка (тофу)»
tcol = ws_lang.max_column + 1
copy_style(ws_lang.cell(1, tcol - 1), ws_lang.cell(1, tcol))
ws_lang.cell(1, tcol).value = 'Тест-строка (тофу)'
ws_lang.column_dimensions[get_column_letter(tcol)].width = 48
written = []  # (код, название по-русски) письменных вариантов в порядке реестра
for r in range(2, ws_lang.max_row + 1):
    code, kind = ws_lang.cell(r, 1).value, ws_lang.cell(r, 2).value
    copy_style(ws_lang.cell(r, tcol - 1), ws_lang.cell(r, tcol))
    if kind == 'письменный':
        written.append((code, ws_lang.cell(r, 3).value))
        ws_lang.cell(r, tcol).value = TOFU_TEST.get(code)
        if code == 'ru-RU':
            # русский — исходный язык, а не перевод: метку «черновик» в приложении не показываем
            ws_lang.cell(r, 10).value = 'ИСХОДНИК'
missing_tofu = [c for c, _ in written if c not in TOFU_TEST]
assert not missing_tofu, missing_tofu
log.append('Перенесены листы ЯЗЫКИ / ИМЕНА / ВЫЧИТКА из таблицы-дополнения v1.0 (22 письменных варианта, голоса, '
           'передача имён, журнал вычитки). В ЯЗЫКИ добавлена колонка «Тест-строка (тофу)». Статус текста ru-RU — '
           'ИСХОДНИК (оригинал, не перевод; метку «черновик» не показываем) — ПОДТВЕРДИТЬ владельцу.')

ws_names = wb['ИМЕНА']
hdr = [ws_names.cell(1, c).value for c in range(1, ws_names.max_column + 1)]
fixed = 0
for r in range(2, ws_names.max_row + 1):
    for c in range(2, ws_names.max_column + 1):
        v = ws_names.cell(r, c).value
        if isinstance(v, str) and 'सेтनेर' in v:
            ws_names.cell(r, c).value = v.replace('सेтनेर', 'सेतनेर')  # кириллическая «т» внутри деванагари
            fixed += 1
    if ws_names.cell(r, 1).value == 'На Чебоксаре (название)':
        ws_names.cell(r, hdr.index('ar') + 1).value = AR_UI['ui.title']
log.append(f'ИМЕНА: исправлена опечатка в хинди/маратхи «Сетнер» (внутри деванагари стояла кириллическая «т»): '
           f'{fixed} ячейки -> सेतनेर (как в листе «Интерфейс»). Добавлено арабское название игры «{AR_UI["ui.title"]}» (ЧЕРНОВИК).')

# ------------------------------------------------------------------ 3. колонки языков в листах текстов
NAME_FIX = {'tt-RU': 'tt'}
written_codes = [c for c, _ in written]
names_ru = dict(written)
order_new = ['ar'] + [c for c in written_codes if c != 'ar']  # арабский — показательная локаль, первой


def lang_cols(ws):
    m = {}
    for c in range(2, ws.max_column + 1):
        h = ws.cell(1, c).value
        if h:
            m[str(h).split('\n')[-1].strip()] = c
    return m


def fix_headers(ws):
    for c in range(2, ws.max_column + 1):
        h = ws.cell(1, c).value
        if not h or '\n' not in str(h):
            continue
        name, code = str(h).rsplit('\n', 1)
        code = NAME_FIX.get(code.strip(), code.strip())
        name = name.replace(' (мордовский)', '')
        ws.cell(1, c).value = f'{name}\n{code}'


added_cols = {}
for sheet, id_col in [('Интерфейс', 1), ('Тексты', 1), ('Легенды', 1)]:
    ws = wb[sheet]
    fix_headers(ws)
    have = lang_cols(ws)
    last = max(have.values())
    width = ws.column_dimensions[get_column_letter(last)].width
    new = [c for c in order_new if c not in have]
    for code in new:
        col = ws.max_column + 1
        copy_style(ws.cell(1, last), ws.cell(1, col))
        title = names_ru[code][:1].upper() + names_ru[code][1:]
        if code == 'ar':
            title += ' · ЧЕРНОВИК (перевод ИИ)'
        ws.cell(1, col).value = f'{title}\n{code}'
        ws.column_dimensions[get_column_letter(col)].width = width
        for r in range(2, ws.max_row + 1):
            copy_style(ws.cell(r, last), ws.cell(r, col))
            if code == 'ar':
                al = copy(ws.cell(r, col).alignment)
                ws.cell(r, col).alignment = Alignment(horizontal='right', vertical=al.vertical or 'top',
                                                      wrap_text=True, readingOrder=2)
    added_cols[sheet] = new

# ------------------------------------------------------------------ 4. новые ключи интерфейса
ws = wb['Интерфейс']
cols = lang_cols(ws)
last_key_row = 1
for r in range(2, ws.max_row + 1):
    v = ws.cell(r, 1).value
    if v and '.' in str(v) and ' ' not in str(v):
        last_key_row = r
existing = {str(ws.cell(r, 1).value) for r in range(2, ws.max_row + 1)}
new_keys = [k for k in NEW_UI if k not in existing]
ws.insert_rows(last_key_row + 1, amount=len(new_keys))
for i, k in enumerate(new_keys):
    r = last_key_row + 1 + i
    for c in range(1, ws.max_column + 1):
        copy_style(ws.cell(2, c), ws.cell(r, c))
    ws.cell(r, 1).value = k
    for lang in L9:
        ws.cell(r, cols[lang]).value = NEW_UI[k][lang]
log.append(f'Интерфейс: +{len(new_keys)} ключей языкового модуля ({", ".join(new_keys)}) на 8 готовых языках + ar. '
           'finds.count и finds.you_found — в формате ICU MessageFormat (ru: one/few/many/other; ar: шесть форм; zh/tr: одна форма).')

# ------------------------------------------------------------------ 5. арабский черновик
ar_ui = ar_t = ar_l = 0
ws = wb['Интерфейс']
c_ar = lang_cols(ws)['ar']
for r in range(2, ws.max_row + 1):
    k = ws.cell(r, 1).value
    if k in AR_UI:
        ws.cell(r, c_ar).value = AR_UI[k]
        ar_ui += 1
    elif k in NEW_UI:
        ar_ui += 1  # уже записан выше
keys_ui = {str(ws.cell(r, 1).value) for r in range(2, ws.max_row + 1) if ws.cell(r, 1).value}
miss_ui = sorted(k for k in keys_ui if (k.split('.')[0] in ('ui', 'pt', 'h', 'q', 'r', 'spk', 'layer', 'today', 'art', 'finds'))
                 and k not in AR_UI and k not in NEW_UI)
assert not miss_ui, f'нет арабского для: {miss_ui}'
ws = wb['Тексты']
c_ar = lang_cols(ws)['ar']
for r in range(2, ws.max_row + 1):
    k = ws.cell(r, 1).value
    if k in AR_TEXTS:
        ws.cell(r, c_ar).value = AR_TEXTS[k]
        ar_t += 1
ws = wb['Легенды']
c_ar = lang_cols(ws)['ar']
for r in range(2, ws.max_row + 1):
    k = ws.cell(r, 1).value
    if k in AR_LEGENDS:
        ws.cell(r, c_ar).value = AR_LEGENDS[k]
        ar_l += 1
assert ar_t == 12 and ar_l == 8, (ar_t, ar_l)
log.append(f'Арабский (MSA, код ar) — показательная RTL-локаль этапа 1: ЧЕРНОВИК, перевод ИИ (Claude), 22.10.2026. '
           f'Интерфейс — {ar_ui} строк, тексты — {ar_t}, предания — {ar_l}. Имена — по листу ИМЕНА. '
           'Вычитка носителем арабского обязательна до статуса выше ЧЕРНОВИК (лист ВЫЧИТКА: «не начато»).')
log.append('Новые пустые колонки для переводчиков (пусто = фолбэк: точный вариант → старший → en-GB → ru-RU): '
           + ', '.join(c for c in added_cols['Интерфейс'] if c != 'ar') + '. Код татарского во всех листах — tt.')
log.append('Генератор content_build.py теперь пишет и locales/<код>.json для каждого письменного варианта листа ЯЗЫКИ '
           '(ключи = ID строк интерфейса, текстов T…, преданий L…), проверяет скобки ICU и печатает отчёт о пропусках. '
           'CSV Unity не менялись; сверка контракта 39/0 сохраняется.')

# ------------------------------------------------------------------ 5б. закон 7 и «тофу» в черновиках переводов
# Чувашская фраза «Турӑ амӑшӗ» (Матерь Божья) в переводах была транслитом латиницей (Tură amăshĕ и т.п.) —
# по закону 7 ТЗ имя пишется родным письмом. Заодно уходит буква ĕ, которой нет в шрифте PT Sans.
# Во французском «XVIIIᵉ» — надстрочная ᵉ, которой нет в шрифте PT Sans: пишем «XVIIIe» (принятое написание).
# Русский оригинал не трогаем (он сверяется с CSV Unity — контракт 39/0); в нём та же фраза набрана с латинскими
# ă и ě вместо чувашских ӑ ӗ — исправить вместе с CSV в следующей версии (решение владельца).
CHUV = 'Турӑ амӑшӗ'
TRANSLIT = ['Tură amăshĕ', 'Tură amăšĕ', 'Tură amăchĕ', 'Tură amăşĕ', 'Tura amashe', 'तुरा अमाश्ये']
law7 = {}
for sheet in ['Интерфейс', 'Тексты', 'Легенды']:
    ws = wb[sheet]
    cols = lang_cols(ws)
    for code, c in cols.items():
        if code == 'ru-RU':
            continue
        for r in range(2, ws.max_row + 1):
            v = ws.cell(r, c).value
            if not isinstance(v, str):
                continue
            nv = v
            for tr in TRANSLIT:
                nv = nv.replace(tr, CHUV)
            if code == 'fr-FR':
                nv = nv.replace('ᵉ', 'e')
            if nv != v:
                ws.cell(r, c).value = nv
                law7.setdefault(code, []).append(str(ws.cell(r, 1).value))
log.append('Закон 7 ТЗ (родное письмо вместо транслита): в черновиках переводов чувашская фраза «Турӑ амӑшӗ» (предание L03) '
           'теперь пишется кириллицей, а не латиницей «Tură amăshĕ» — ' + '; '.join(f'{k}: {", ".join(sorted(set(v)))}' for k, v in sorted(law7.items()))
           + '. Во французском «XVIIIᵉ» → «XVIIIe» (надстрочной ᵉ нет в шрифте). Русский оригинал не менялся: в нём фраза набрана '
           'латинскими ă/ě вместо чувашских ӑ/ӗ — исправить вместе с CSV Unity (ПРЕДЛОЖЕНИЕ владельцу).')

# ------------------------------------------------------------------ 6. журнал
w2 = wb['Как устроено']
w2.cell(2, 2).value = 'Краеведческий квест · проект «Спутник гостя Чебоксар» · версия 2.7 · 22 октября 2026'
r = w2.max_row + 2
w2.cell(r, 2).value = 'ЖУРНАЛ ИЗМЕНЕНИЙ v2.6 → v2.7 (22.10.2026, языковой модуль — этап 1 «Каркас», приложение 1.1.0)'
r += 1
for line in log:
    w2.cell(r, 3).value = line
    r += 1
w2.cell(r, 3).value = ('Статусы: всё новое в этой версии — ЧЕРНОВИК. В листе ВЫЧИТКА все локали — «не начато». '
                       'В приложении варианты со статусом ЧЕРНОВИК помечены служебной меткой ui.lang_draft_badge.')

# порядок листов: служебные языковые — после «Интерфейс»
order = [s.title for s in wb.worksheets]
for name in ['ВЫЧИТКА', 'ИМЕНА', 'ЯЗЫКИ']:
    ws = wb[name]
    wb.move_sheet(ws, offset=order.index('Интерфейс') + 1 - wb.index(ws))
    order = [s.title for s in wb.worksheets]

wb.save(DST)
print('Сохранено:', DST)
print('Листы:', [s.title for s in wb.worksheets])
print('Новые колонки:', added_cols)
print('Новые ключи:', new_keys)
print(f'ar: интерфейс {ar_ui}, тексты {ar_t}, предания {ar_l}')
