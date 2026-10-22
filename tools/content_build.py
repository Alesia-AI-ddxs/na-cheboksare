# -*- coding: utf-8 -*-
"""Генератор контента приложения «На Чебоксаре» из мастер-таблицы КОНТЕНТ_На_Чебоксаре.
Правило проекта: приложение ничего не выдумывает — все тексты берутся из таблицы.
Структура точек повторяет Story.cs игры Unity v0.3 и расширена на точки 4–8.

Запуск:  python3 -I content_build.py <путь_к_xlsx> <выход_content.json>
"""
import sys, json, re
import openpyxl

XLSX = sys.argv[1]
OUT = sys.argv[2]

LANG_CODES = ['ru-RU', 'en-GB', 'en-US', 'de-DE', 'fr-FR', 'zh-Hans', 'hi-IN', 'tr-TR', 'cv', 'tt-RU', 'myv', 'mdf']
READY = {'ru-RU', 'en-GB', 'en-US', 'de-DE', 'fr-FR', 'zh-Hans', 'hi-IN', 'tr-TR'}

# ---------------------------------------------------------------- структура маршрута (= Story.cs + точки 4–8)
def beat(key, speaker, layer, era):
    return {'key': key, 'speaker': speaker, 'layer': layer, 'era': era}

POINTS = [
    {'id': 0, 'stage': 1, 'scene': 'night', 'era': '1469', 'nameKey': 'pt.0',
     'beats': [beat('T00_PROLOG', 'spk.chronicler', 'doc', '1469')],
     'legends': [beat('L02_GOD_1469', 'spk.bookman', 'leg', '1469')],
     'quiz': None, 'artifacts': ['A1_KRESALO'], 'hintKey': None, 'todayKey': None},
    {'id': 1, 'stage': 1, 'scene': 'pier', 'era': '1636', 'nameKey': 'pt.1',
     'beats': [beat('T01_OLEARIUS', 'spk.olearius', 'doc', '1636')],
     'legends': [],
     'quiz': {'q': 'q.1', 'r': 'r.1', 'options': ['30', '300', '3000'], 'correct': 1},
     'artifacts': ['A2_VESLO'], 'hintKey': 'h.1', 'todayKey': 'today.1'},
    {'id': 2, 'stage': 1, 'scene': 'kremlin', 'era': '1555', 'nameKey': 'pt.2',
     'beats': [beat('T02_NAGOY', 'spk.chronicler', 'doc', '1555'),
               beat('T03_ROSPIS_1613', 'spk.clerk', 'doc', '1613'),
               beat('T04_KREMLIN_NARR', 'spk.setner', 'sci', '1613')],
     'legends': [beat('L01_VYRODKOV', 'spk.strelets', 'leg', '1613'),
                 beat('L05_MERA', 'spk.strelets', 'leg', '1613'),
                 beat('L08_VOEVODA', 'spk.bookman', 'leg', '1555')],
     'quiz': {'q': 'q.2', 'r': 'r.2', 'options': ['3', '5', '9'], 'correct': 1},
     'artifacts': ['A3_BERESTA', 'A4_ZERKALO', 'A5_BERDYSH'], 'hintKey': 'h.2', 'todayKey': 'today.2'},
    {'id': 3, 'stage': 1, 'scene': 'cathedral', 'era': '1650', 'nameKey': 'pt.3',
     'beats': [beat('T05_SOBOR_NARR', 'spk.setner', 'sci', '1650')],
     'legends': [beat('L03_TORAMUSH', 'spk.parishioner', 'leg', '1650'),
                 beat('L04_POZHAR', 'spk.bookman', 'leg', '1650'),
                 beat('L06_GOD_SOBORA', 'spk.bookman', 'leg', '1650'),
                 beat('L07_FRESKI', 'spk.bookman', 'leg', '1650')],
     'quiz': {'q': 'q.3', 'r': 'r.3', 'options': ['7', '11', '12'], 'correct': 1},
     'artifacts': ['A6_IZRAZETS'], 'hintKey': 'h.3', 'todayKey': 'today.3'},
    {'id': 8, 'stage': 1, 'scene': 'bay', 'era': 'today', 'nameKey': 'pt.8',
     'beats': [beat('T11_FIGURKI_NARR', 'spk.setner', 'sci', 'today')],
     'legends': [],
     'quiz': {'q': 'q.8', 'r': 'r.8', 'options': ['4', '6', '12'], 'correct': 1},
     'artifacts': ['A12_FIGURKI'], 'hintKey': 'h.8', 'todayKey': 'today.8', 'stageFinish': 'A7_KOPEYKA'},
    # ---- этап II «Купеческий город»
    {'id': 4, 'stage': 2, 'scene': 'mansion', 'era': '1900', 'nameKey': 'pt.4',
     'beats': [beat('T06_EFREMOV_NARR', 'spk.setner', 'sci', '1900')],
     'legends': [],
     'quiz': {'q': 'q.4', 'r': 'r.4', 'options': ['3 000', '30 000', '300 000'], 'correct': 1},
     'artifacts': ['A8_TORGOVY_DOM'], 'hintKey': 'h.4', 'todayKey': 'today.4'},
    {'id': 5, 'stage': 2, 'scene': 'granary', 'era': '1773', 'nameKey': 'pt.5',
     'beats': [beat('T07_FONBRANT', 'spk.fonbrant', 'doc', '1773'),
               beat('T08_AMBARY_NARR', 'spk.setner', 'sci', '1773')],
     'legends': [],
     'quiz': {'q': 'q.5', 'r': 'r.5', 'options': ['9', '49', '490'], 'correct': 1},
     'artifacts': ['A9_DOKLAD_1773'], 'hintKey': 'h.5', 'todayKey': 'today.5'},
    {'id': 6, 'stage': 2, 'scene': 'school', 'era': '1904', 'nameKey': 'pt.6',
     'beats': [beat('T09_GIMNAZIA_NARR', 'spk.setner', 'sci', '1904')],
     'legends': [],
     'quiz': {'q': 'q.6', 'r': 'r.6', 'options': ['46', '146', '460'], 'correct': 0},
     'artifacts': ['A10_PORTRETY_1908'], 'hintKey': 'h.6', 'todayKey': 'today.6'},
    {'id': 7, 'stage': 2, 'scene': 'hill', 'era': '1899', 'nameKey': 'pt.7',
     'beats': [beat('T10_GORKA_NARR', 'spk.setner', 'sci', '1899')],
     'legends': [],
     'quiz': {'q': 'q.7', 'r': 'r.7', 'options': ['1961', '1991', '2011'], 'correct': 1},
     'artifacts': ['A11_PROEKT_1907'], 'hintKey': 'h.7', 'todayKey': 'today.7', 'stageFinish': None},
]
STAGES = [
    {'id': 1, 'nameKey': 'ui.map', 'points': [0, 1, 2, 3, 8], 'finishArtifact': 'A7_KOPEYKA'},
    {'id': 2, 'nameKey': 'ui.map2', 'points': [4, 5, 6, 7], 'finishArtifact': None},
]
# тип находки: museum — подлинный предмет из раскопок; evidence — документальное свидетельство; game — игровая
ARTIFACT_KIND = {
    'A1_KRESALO': 'museum', 'A2_VESLO': 'museum', 'A3_BERESTA': 'museum', 'A4_ZERKALO': 'museum',
    'A5_BERDYSH': 'museum', 'A6_IZRAZETS': 'museum', 'A7_KOPEYKA': 'museum',
    'A8_TORGOVY_DOM': 'evidence', 'A9_DOKLAD_1773': 'evidence', 'A10_PORTRETY_1908': 'evidence',
    'A11_PROEKT_1907': 'evidence', 'A12_FIGURKI': 'game',
}
ARTIFACT_ICON = {
    'A1_KRESALO': 'firesteel', 'A2_VESLO': 'oar', 'A3_BERESTA': 'bark', 'A4_ZERKALO': 'mirror',
    'A5_BERDYSH': 'poleaxe', 'A6_IZRAZETS': 'tile', 'A7_KOPEYKA': 'coin',
    'A8_TORGOVY_DOM': 'scroll', 'A9_DOKLAD_1773': 'scroll', 'A10_PORTRETY_1908': 'portrait',
    'A11_PROEKT_1907': 'blueprint', 'A12_FIGURKI': 'figurine',
}

# ---------------------------------------------------------------- чтение таблицы
wb = openpyxl.load_workbook(XLSX, data_only=True)

def header_lang_map(ws, first_lang_col):
    """колонка -> код языка, по коду в заголовке (последняя строка заголовка)"""
    m = {}
    for c in range(first_lang_col, ws.max_column + 1):
        h = ws.cell(1, c).value
        if not h:
            continue
        code = str(h).strip().split('\n')[-1].strip()
        if code in LANG_CODES:
            m[c] = code
    return m

def clean(v):
    if v is None:
        return ''
    s = str(v).replace('\r\n', '\n').strip()
    return s

warnings = []

# --- Интерфейс
ws = wb['Интерфейс']
lm = header_lang_map(ws, 2)
ui = {}
for r in range(2, ws.max_row + 1):
    key = clean(ws.cell(r, 1).value)
    if not key or not re.match(r'^[a-z]+\.', key):
        continue
    ui[key] = {code: clean(ws.cell(r, c).value) for c, code in lm.items() if clean(ws.cell(r, c).value)}

# --- Тексты
ws = wb['Тексты']
lm = header_lang_map(ws, 9)
texts = {}
for r in range(2, ws.max_row + 1):
    tid = clean(ws.cell(r, 1).value)
    if not tid:
        continue
    texts[tid] = {
        'point': clean(ws.cell(r, 2).value),
        'who': clean(ws.cell(r, 3).value),
        'voice': clean(ws.cell(r, 4).value),
        'layer': clean(ws.cell(r, 5).value),
        'source': clean(ws.cell(r, 6).value),
        'status': clean(ws.cell(r, 7).value),
        'text': {code: clean(ws.cell(r, c).value) for c, code in lm.items() if clean(ws.cell(r, c).value)},
    }

# --- Легенды
ws = wb['Легенды']
lm = header_lang_map(ws, 6)
legends = {}
for r in range(2, ws.max_row + 1):
    lid = clean(ws.cell(r, 1).value)
    if not lid:
        continue
    legends[lid] = {
        'point': clean(ws.cell(r, 2).value),
        'who': clean(ws.cell(r, 3).value),
        'layer': clean(ws.cell(r, 4).value),
        'dispute': clean(ws.cell(r, 5).value),
        'text': {code: clean(ws.cell(r, c).value) for c, code in lm.items() if clean(ws.cell(r, c).value)},
    }

# --- Факты (только русский; это справочный слой)
ws = wb['Факты']
facts = []
for r in range(2, ws.max_row + 1):
    if ws.cell(r, 3).value is None:
        continue
    facts.append({
        'n': ws.cell(r, 1).value, 'point': clean(ws.cell(r, 2).value), 'topic': clean(ws.cell(r, 3).value),
        'year': clean(ws.cell(r, 4).value), 'text': clean(ws.cell(r, 5).value), 'layer': clean(ws.cell(r, 6).value),
        'source': clean(ws.cell(r, 7).value), 'link': clean(ws.cell(r, 8).value),
    })

# --- Источники
ws = wb['Источники']
sources = []
for r in range(2, ws.max_row + 1):
    if ws.cell(r, 2).value is None:
        continue
    sources.append({
        'n': ws.cell(r, 1).value, 'title': clean(ws.cell(r, 2).value), 'gives': clean(ws.cell(r, 3).value),
        'link': clean(ws.cell(r, 4).value), 'status': clean(ws.cell(r, 5).value), 'todo': clean(ws.cell(r, 6).value),
    })

# --- Языки
ws = wb['Языки']
languages = []
for r in range(2, ws.max_row + 1):
    code = clean(ws.cell(r, 1).value)
    if code in LANG_CODES:
        languages.append({'code': code, 'name': clean(ws.cell(r, 2).value), 'ready': code in READY})

# --- Как устроено: версия таблицы
ws = wb['Как устроено']
table_version = clean(ws.cell(2, 2).value)

# ---------------------------------------------------------------- сборка точек с проверкой
artifacts = {}
for aid, kind in ARTIFACT_KIND.items():
    nk, dk = f'art.{aid}.name', f'art.{aid}.desc'
    if nk not in ui:
        warnings.append(f'нет {nk} в Интерфейсе')
    artifacts[aid] = {'id': aid, 'kind': kind, 'icon': ARTIFACT_ICON[aid], 'nameKey': nk, 'descKey': dk}

points_out = []
for p in POINTS:
    pid = p['id']
    out = dict(p)
    out['facts'] = [f['n'] for f in facts if f['point'] == str(pid)]
    missing = []
    for b in p['beats']:
        if b['key'] not in texts:
            missing.append(b['key'])
        if b['speaker'] not in ui:
            missing.append(b['speaker'])
    for b in p['legends']:
        if b['key'] not in legends:
            missing.append(b['key'])
    if p['quiz']:
        for k in (p['quiz']['q'], p['quiz']['r']):
            if k not in ui:
                missing.append(k)
    for k in (p['hintKey'], p['todayKey'], p['nameKey']):
        if k and k not in ui:
            missing.append(k)
    out['ready'] = not missing
    out['missing'] = missing
    if missing:
        warnings.append(f'точка {pid}: не хватает {missing}')
    points_out.append(out)

content = {
    'app': 'На Чебоксаре · Туризм — Чебоксары',
    'contentVersion': table_version,
    'generatedFrom': XLSX.split('/')[-1],
    'languages': languages,
    'stages': STAGES,
    'points': points_out,
    'texts': texts,
    'legends': legends,
    'artifacts': artifacts,
    'ui': ui,
    'facts': facts,
    'sources': sources,
    'links': {
        'setnerTelegram': 'https://t.me/Setner_Cheboksary_bot',
        'web': 'https://varya.esteticdent21.ru/turizm/',
        'webGame': 'https://varya.esteticdent21.ru/turizm/igra/',
    },
}
with open(OUT, 'w', encoding='utf-8') as f:
    json.dump(content, f, ensure_ascii=False, indent=1)

print('content.json:', OUT)
print('языков:', len(languages), '| ui-ключей:', len(ui), '| текстов:', len(texts), '| преданий:', len(legends),
      '| фактов:', len(facts), '| источников:', len(sources))
print('точки готовы:', [p['id'] for p in points_out if p['ready']], '| не готовы:', [p['id'] for p in points_out if not p['ready']])
for w in warnings:
    print('ПРЕДУПРЕЖДЕНИЕ:', w)
